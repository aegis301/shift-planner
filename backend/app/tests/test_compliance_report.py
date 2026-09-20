from datetime import UTC, date, datetime
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.deps import get_db
from app.core.security import hash_password
from app.main import app
from app.models import Account, Organization, ShiftGroup, User, UserShiftGroup
from app.models.base import Base
from app.services.authz import ROLE_PLANNER
from app.services.compliance_report import build_compliance_report
from app.services.exports import export_compliance_report_pdf, export_compliance_report_xlsx
from app.services.rules import build_plan_state, evaluate_plan_state
from app.services.validation import filter_warnings_for_shift_group

RULES = [
    {
        "type": "max_daily_working_time",
        "severity": "warning",
        "base_hours": "10",
        "extended_hours": "12",
        "extension_requires_duty_hours": "24",
    },
    {
        "type": "weekly_average_cap",
        "severity": "warning",
        "hours": "48",
        "reference_period_months": 6,
        "rolling": True,
    },
    {
        "type": "opt_out_weekly_cap",
        "severity": "warning",
        "hours_by_tier": {"standard": "48", "stufe_i": "54"},
        "reference_period_months": 6,
    },
    {"type": "max_consecutive_work_days", "severity": "warning", "days": 6},
    {
        "type": "max_duties_per_period",
        "severity": "warning",
        "count": 4,
        "period": "month",
        "additional_allowance_per_quarter": 0,
    },
    {
        "type": "documentation_requirement",
        "severity": "info",
        "threshold_hours": "8",
        "retention_months": 24,
    },
]


def _seed_membership(db, email: str, password: str, org_id: int, role: str) -> User:
    acc = Account(email=email.lower(), hashed_password=hash_password(password))
    db.add(acc)
    db.flush()
    user = User(account_id=acc.id, organization_id=org_id, role=role, locale="de")
    db.add(user)
    return user


@pytest.fixture()
def client():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    testing_session = sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False)
    Base.metadata.create_all(engine)
    with testing_session() as db:
        db.add(Organization(id=1, name="Default", slug="default", plan_tier="team"))
        db.flush()
        _seed_membership(db, "admin@example.com", "secret", 1, "admin")
        planner = _seed_membership(db, "planner@example.com", "plannersecret", 1, ROLE_PLANNER)
        db.flush()
        group = ShiftGroup(organization_id=1, code="sg1", name="SG1", display_order=0)
        other = ShiftGroup(organization_id=1, code="sg2", name="SG2", display_order=1)
        db.add(group)
        db.add(other)
        db.flush()
        db.add(UserShiftGroup(user_id=planner.id, shift_group_id=group.id))
        db.commit()

    def override_get_db():
        db = testing_session()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client, testing_session
    app.dependency_overrides.clear()


def login_admin(test_client: TestClient) -> None:
    assert (
        test_client.post(
            "/api/v1/auth/login",
            json={"email": "admin@example.com", "password": "secret", "organization_slug": "default"},
        ).status_code
        == 200
    )


def login_planner(test_client: TestClient) -> None:
    assert (
        test_client.post(
            "/api/v1/auth/login",
            json={
                "email": "planner@example.com",
                "password": "plannersecret",
                "organization_slug": "default",
            },
        ).status_code
        == 200
    )


def _fixture_month(test_client: TestClient) -> tuple[int, int, int]:
    login_admin(test_client)
    member = test_client.post(
        "/api/v1/team-members",
        json={
            "first_name": "Pat",
            "last_name": "Duty",
            "email": "pat.duty@example.com",
            "employment_percentage": 100,
        },
    )
    assert member.status_code == 200, member.text
    member_id = member.json()["id"]
    assert (
        test_client.put(
            "/api/v1/shift-groups/1/memberships",
            json={"memberships": [{"team_member_id": member_id, "start_date": "2026-01-01", "end_date": None}]},
        ).status_code
        == 200
    )
    template = test_client.post(
        "/api/v1/shift-templates",
        json={"code": "BD", "name": "BD", "category": "bereitschaftsdienst"},
    )
    assert template.status_code == 200, template.text
    template_id = template.json()["id"]
    assert (
        test_client.post(
            f"/api/v1/shift-templates/{template_id}/variants",
            json={
                "label": "24h",
                "start_day_class": "any",
                "starts_at": "08:00:00",
                "ends_at": "08:00:00",
                "end_day_offset": 1,
                "required_count": 1,
            },
        ).status_code
        == 200
    )
    assert (
        test_client.put("/api/v1/shift-groups/1/shift-templates", json={"shift_template_ids": [template_id]}).status_code
        == 200
    )
    period = test_client.post("/api/v1/planning-periods", json={"year": 2026, "month": 3})
    assert period.status_code == 200, period.text
    period_id = period.json()["id"]
    roster = test_client.get(f"/api/v1/roster-matrix/{period_id}?shift_group_id=1").json()
    slot = next(item for item in roster["slots"] if item["slot_date"] == "2026-03-02")
    assigned = test_client.put(
        "/api/v1/roster-matrix/assignments",
        json={"roster_slot_id": slot["id"], "team_member_id": member_id},
    )
    assert assigned.status_code == 200, assigned.text
    created = test_client.post(
        "/api/v1/work-time-rule-sets",
        json={"name": "ArbZG-Test", "is_active": True, "rules": RULES},
    )
    assert created.status_code == 201, created.text
    return period_id, member_id, created.json()["id"]


def test_report_matches_engine_findings_for_fixture_month(client):
    test_client, testing_session = client
    period_id, member_id, rule_set_id = _fixture_month(test_client)
    db = testing_session()
    try:
        report = build_compliance_report(db, period_id, organization_id=1, shift_group_id=1)
        state = build_plan_state(
            db,
            organization_id=1,
            start_date=date(2026, 3, 1),
            end_date=date(2026, 3, 31),
        )
        engine = filter_warnings_for_shift_group(
            db,
            evaluate_plan_state(state, db=db),
            planning_period_id=period_id,
            organization_id=1,
            shift_group_id=1,
        )
        report_keys = {(row.code, row.team_member_id, row.date, row.message) for row in report.findings}
        engine_keys = {(row.code, row.team_member_id, row.date, row.message) for row in engine}
        assert report_keys == engine_keys
        assert any(row.code == "WORKTIME_MAX_DAILY" and row.team_member_id == member_id for row in report.findings)
        member = next(row for row in report.members if row.team_member_id == member_id)
        assert member.statutory_minutes > 0
        assert member.credited_minutes >= 0
        assert report.rule_set is not None
        assert report.rule_set.id == rule_set_id
        assert report.rule_set.version == 1
        assert report.rule_set.name == "ArbZG-Test"
    finally:
        db.close()


def test_report_is_deterministic_on_rerun(client):
    test_client, testing_session = client
    period_id, _member_id, _rule_set_id = _fixture_month(test_client)
    stamp = datetime(2026, 4, 1, 12, 0, tzinfo=UTC)
    db = testing_session()
    try:
        first = build_compliance_report(
            db, period_id, organization_id=1, shift_group_id=1, generated_at=stamp
        )
        second = build_compliance_report(
            db, period_id, organization_id=1, shift_group_id=1, generated_at=stamp
        )
        assert first.model_dump(mode="json") == second.model_dump(mode="json")
        xlsx_a = export_compliance_report_xlsx(
            db, period_id, organization_id=1, shift_group_id=1, generated_at=stamp
        )
        xlsx_b = export_compliance_report_xlsx(
            db, period_id, organization_id=1, shift_group_id=1, generated_at=stamp
        )
        sheet_a = load_workbook(BytesIO(xlsx_a)).active
        sheet_b = load_workbook(BytesIO(xlsx_b)).active
        values_a = [[cell.value for cell in row] for row in sheet_a.iter_rows()]
        values_b = [[cell.value for cell in row] for row in sheet_b.iter_rows()]
        assert values_a == values_b
    finally:
        db.close()


def test_exports_print_rule_set_version(client):
    test_client, testing_session = client
    period_id, _member_id, _rule_set_id = _fixture_month(test_client)
    stamp = datetime(2026, 4, 1, 12, 0, tzinfo=UTC)
    login_admin(test_client)
    xlsx = test_client.get(f"/api/v1/exports/compliance-report/{period_id}.xlsx?shift_group_id=1")
    assert xlsx.status_code == 200, xlsx.text
    assert xlsx.content.startswith(b"PK")
    db = testing_session()
    try:
        body = export_compliance_report_xlsx(
            db, period_id, organization_id=1, shift_group_id=1, generated_at=stamp
        )
        sheet = load_workbook(BytesIO(body)).active
        assert "ArbZG-Test" in str(sheet["A3"].value)
        assert "v1" in str(sheet["A3"].value)
        assert "2026-04-01T12:00:00+00:00" in str(sheet["A4"].value)
        pdf_body = export_compliance_report_pdf(
            db, period_id, organization_id=1, shift_group_id=1, generated_at=stamp
        )
        assert b"ArbZG-Test" in pdf_body
        assert b"v1" in pdf_body
        assert b"Generated:" in pdf_body
    finally:
        db.close()
    pdf = test_client.get(f"/api/v1/exports/compliance-report/{period_id}.pdf?shift_group_id=1")
    assert pdf.status_code == 200, pdf.text
    assert pdf.content.startswith(b"%PDF")
    assert b"ArbZG-Test" in pdf.content
    assert b"v1" in pdf.content


def test_planner_scope_enforced_on_compliance_report(client):
    test_client, _session = client
    period_id, _member_id, _rule_set_id = _fixture_month(test_client)
    login_planner(test_client)
    missing = test_client.get(f"/api/v1/compliance-report/{period_id}")
    assert missing.status_code == 403
    other = test_client.get(f"/api/v1/compliance-report/{period_id}?shift_group_id=2")
    assert other.status_code == 403
    allowed = test_client.get(f"/api/v1/compliance-report/{period_id}?shift_group_id=1")
    assert allowed.status_code == 200, allowed.text
    assert allowed.json()["rule_set"]["version"] == 1
    denied_xlsx = test_client.get(f"/api/v1/exports/compliance-report/{period_id}.xlsx")
    assert denied_xlsx.status_code == 403
    ok_xlsx = test_client.get(f"/api/v1/exports/compliance-report/{period_id}.xlsx?shift_group_id=1")
    assert ok_xlsx.status_code == 200
