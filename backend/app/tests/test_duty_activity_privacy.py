from datetime import date
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.deps import get_db
from app.core.security import hash_password
from app.main import app
from app.models import (
    Account,
    AuditLog,
    Organization,
    ShiftGroup,
    TeamMember,
    TimeEntry,
    User,
    UserShiftGroup,
)
from app.models.base import Base
from app.services.duty_activity_privacy import purge_expired_duty_activity_episodes


def _seed_membership(db, email: str, password: str, org_id: int, role: str) -> User:
    acc = Account(email=email.lower(), hashed_password=hash_password(password))
    db.add(acc)
    db.flush()
    user = User(account_id=acc.id, organization_id=org_id, role=role, locale="de")
    db.add(user)
    return user


def _client_session():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    testing_session = sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False)
    Base.metadata.create_all(engine)
    with testing_session() as db:
        db.add(Organization(id=1, name="Default", slug="default", plan_tier="team"))
        db.flush()
        _seed_membership(db, "admin@example.com", "secret", 1, "admin")
        planner = _seed_membership(db, "planner@example.com", "plannersecret", 1, "planner")
        portal = _seed_membership(db, "doc@example.com", "docsecret", 1, "team_member")
        db.flush()
        member = TeamMember(
            organization_id=1,
            first_name="UniqueAlpha",
            last_name="UniqueBeta",
            email="uniquealpha.uniquebeta@example.com",
            user_id=portal.id,
        )
        db.add(member)
        sg = ShiftGroup(organization_id=1, code="default_sg", name="Default SG", display_order=0)
        db.add(sg)
        db.flush()
        db.add(UserShiftGroup(user_id=planner.id, shift_group_id=sg.id))
        db.commit()

    def override_get_db():
        db = testing_session()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client, testing_session
    app.dependency_overrides.clear()
    engine.dispose()


def login_admin(client: TestClient) -> None:
    assert client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.com", "password": "secret", "organization_slug": "default"},
    ).status_code == 200


def login_planner(client: TestClient) -> None:
    assert client.post(
        "/api/v1/auth/login",
        json={"email": "planner@example.com", "password": "plannersecret", "organization_slug": "default"},
    ).status_code == 200


def login_team_member(client: TestClient) -> None:
    assert client.post(
        "/api/v1/auth/login",
        json={"email": "doc@example.com", "password": "docsecret", "organization_slug": "default"},
    ).status_code == 200


def set_shift_group_membership(client: TestClient, *, team_member_id: int) -> None:
    response = client.put(
        "/api/v1/shift-groups/1/memberships",
        json={"memberships": [{"team_member_id": team_member_id, "start_date": "2026-01-01", "end_date": None}]},
    )
    assert response.status_code == 200


def create_member(client: TestClient, *, first_name: str, last_name: str, email: str) -> int:
    response = client.post(
        "/api/v1/team-members",
        json={
            "first_name": first_name,
            "last_name": last_name,
            "email": email,
            "employment_percentage": 100,
        },
    )
    assert response.status_code == 200
    member_id = response.json()["id"]
    set_shift_group_membership(client, team_member_id=member_id)
    return member_id


def _create_category_template(client: TestClient, *, code: str, category: str) -> int:
    template = client.post("/api/v1/shift-templates", json={"code": code, "name": code, "category": category})
    assert template.status_code == 200
    template_id = template.json()["id"]
    assert (
        client.post(
            f"/api/v1/shift-templates/{template_id}/variants",
            json={
                "label": code,
                "start_day_class": "any",
                "starts_at": "08:00:00",
                "ends_at": "08:00:00",
                "end_day_offset": 1,
                "required_count": 1,
            },
        ).status_code
        == 200
    )
    return template_id


def _slot_for_day(client: TestClient, period_id: int, slot_date: str) -> dict:
    roster = client.get(f"/api/v1/roster-matrix/{period_id}").json()
    return next(item for item in roster["slots"] if item["slot_date"] == slot_date)


def test_planner_individual_episodes_denied_aggregates_ok():
    for test_client, _session in _client_session():
        login_admin(test_client)
        own_id = None
        login_team_member(test_client)
        own_id = test_client.get("/api/v1/auth/me").json()["team_member_id"]
        login_admin(test_client)
        set_shift_group_membership(test_client, team_member_id=own_id)
        _create_category_template(test_client, code="BD", category="bereitschaftsdienst")
        period = test_client.post("/api/v1/planning-periods", json={"year": 2026, "month": 3})
        assert period.status_code == 200
        period_id = period.json()["id"]
        slot = _slot_for_day(test_client, period_id, "2026-03-02")
        assigned = test_client.put(
            "/api/v1/roster-matrix/assignments",
            json={"roster_slot_id": slot["id"], "team_member_id": own_id},
        )
        assert assigned.status_code == 200
        login_team_member(test_client)
        assert test_client.post("/api/v1/duty-activity/purpose/acknowledge").status_code == 200
        created = test_client.post(
            "/api/v1/duty-activity",
            json={
                "roster_slot_id": slot["id"],
                "kind": "in_duty_activity",
                "started_at": "2026-03-02T08:00:00+00:00",
                "ended_at": "2026-03-02T10:00:00+00:00",
                "reason": {"code": "ward"},
            },
        )
        assert created.status_code == 200, created.text
        login_planner(test_client)
        denied = test_client.get(f"/api/v1/duty-activity?team_member_id={own_id}")
        assert denied.status_code == 403
        util = test_client.get(f"/api/v1/duty-activity/utilization/{period_id}?shift_group_id=1")
        assert util.status_code == 200, util.text
        login_admin(test_client)
        admin_denied = test_client.get(f"/api/v1/duty-activity?team_member_id={own_id}")
        assert admin_denied.status_code == 403


def test_granting_individual_access_is_audited_then_planner_can_read():
    for test_client, testing_session in _client_session():
        login_team_member(test_client)
        own_id = test_client.get("/api/v1/auth/me").json()["team_member_id"]
        login_admin(test_client)
        set_shift_group_membership(test_client, team_member_id=own_id)
        _create_category_template(test_client, code="BD", category="bereitschaftsdienst")
        period_id = test_client.post("/api/v1/planning-periods", json={"year": 2026, "month": 3}).json()["id"]
        slot = _slot_for_day(test_client, period_id, "2026-03-02")
        assert (
            test_client.put(
                "/api/v1/roster-matrix/assignments",
                json={"roster_slot_id": slot["id"], "team_member_id": own_id},
            ).status_code
            == 200
        )
        login_team_member(test_client)
        assert test_client.post("/api/v1/duty-activity/purpose/acknowledge").status_code == 200
        assert (
            test_client.post(
                "/api/v1/duty-activity",
                json={
                    "roster_slot_id": slot["id"],
                    "kind": "in_duty_activity",
                    "started_at": "2026-03-02T08:00:00+00:00",
                    "ended_at": "2026-03-02T10:00:00+00:00",
                },
            ).status_code
            == 200
        )
        login_admin(test_client)
        patched = test_client.patch(
            "/api/v1/organization/duty-activity-access-policy",
            json={"individual_read_roles": ["planner"]},
        )
        assert patched.status_code == 200, patched.text
        assert patched.json()["individual_read_roles"] == ["planner"]
        with testing_session() as db:
            logs = list(
                db.scalars(
                    select(AuditLog).where(
                        AuditLog.entity_type == "duty_activity_access_policy",
                        AuditLog.action == "update",
                    )
                )
            )
            assert len(logs) == 1
            assert logs[0].details["individual_read_roles"] == ["planner"]
        login_planner(test_client)
        listed = test_client.get(f"/api/v1/duty-activity?team_member_id={own_id}")
        assert listed.status_code == 200, listed.text
        assert len(listed.json()) == 1
        with testing_session() as db:
            reads = list(
                db.scalars(
                    select(AuditLog).where(
                        AuditLog.entity_type == "duty_activity",
                        AuditLog.action == "read",
                    )
                )
            )
            assert any(row.entity_id == str(own_id) for row in reads)


def test_purpose_must_be_acknowledged_before_rest_capture():
    for test_client, _session in _client_session():
        login_team_member(test_client)
        own_id = test_client.get("/api/v1/auth/me").json()["team_member_id"]
        login_admin(test_client)
        set_shift_group_membership(test_client, team_member_id=own_id)
        _create_category_template(test_client, code="BD", category="bereitschaftsdienst")
        period_id = test_client.post("/api/v1/planning-periods", json={"year": 2026, "month": 3}).json()["id"]
        slot = _slot_for_day(test_client, period_id, "2026-03-02")
        assert (
            test_client.put(
                "/api/v1/roster-matrix/assignments",
                json={"roster_slot_id": slot["id"], "team_member_id": own_id},
            ).status_code
            == 200
        )
        login_admin(test_client)
        assert (
            test_client.patch(
                "/api/v1/organization/duty-activity-access-policy",
                json={"purpose_statement": "Activity is recorded for utilization only."},
            ).status_code
            == 200
        )
        login_team_member(test_client)
        purpose = test_client.get("/api/v1/duty-activity/purpose")
        assert purpose.status_code == 200
        assert purpose.json()["acknowledged"] is False
        assert "utilization" in purpose.json()["purpose_statement"]
        blocked = test_client.post(
            "/api/v1/duty-activity",
            json={
                "roster_slot_id": slot["id"],
                "kind": "in_duty_activity",
                "started_at": "2026-03-02T08:00:00+00:00",
                "ended_at": "2026-03-02T10:00:00+00:00",
            },
        )
        assert blocked.status_code == 403
        ack = test_client.post("/api/v1/duty-activity/purpose/acknowledge")
        assert ack.status_code == 200
        assert ack.json()["acknowledged"] is True
        created = test_client.post(
            "/api/v1/duty-activity",
            json={
                "roster_slot_id": slot["id"],
                "kind": "in_duty_activity",
                "started_at": "2026-03-02T08:00:00+00:00",
                "ended_at": "2026-03-02T10:00:00+00:00",
            },
        )
        assert created.status_code == 200, created.text


def test_works_council_export_has_no_individual_data_and_suppresses_small_groups():
    for test_client, _session in _client_session():
        login_admin(test_client)
        first_id = None
        login_team_member(test_client)
        first_id = test_client.get("/api/v1/auth/me").json()["team_member_id"]
        login_admin(test_client)
        set_shift_group_membership(test_client, team_member_id=first_id)
        second_id = create_member(
            test_client,
            first_name="GammaPerson",
            last_name="DeltaPerson",
            email="gammaperson.deltaperson@example.com",
        )
        _create_category_template(test_client, code="BD", category="bereitschaftsdienst")
        period_id = test_client.post("/api/v1/planning-periods", json={"year": 2026, "month": 3}).json()["id"]
        slot_a = _slot_for_day(test_client, period_id, "2026-03-02")
        slot_b = _slot_for_day(test_client, period_id, "2026-03-03")
        assert (
            test_client.put(
                "/api/v1/roster-matrix/assignments",
                json={"roster_slot_id": slot_a["id"], "team_member_id": first_id},
            ).status_code
            == 200
        )
        assert (
            test_client.put(
                "/api/v1/roster-matrix/assignments",
                json={"roster_slot_id": slot_b["id"], "team_member_id": second_id},
            ).status_code
            == 200
        )
        small = test_client.get(f"/api/v1/exports/duty-activity/works-council/{period_id}.xlsx")
        assert small.status_code == 200, small.text
        workbook = load_workbook(BytesIO(small.content))
        values = [
            str(value)
            for row in workbook.active.iter_rows(values_only=True)
            for value in row
            if value is not None
        ]
        blob = " ".join(values).lower()
        assert "uniquealpha" not in blob
        assert "uniquebeta" not in blob
        assert "gammaperson" not in blob
        assert "deltaperson" not in blob
        assert "uniquealpha.uniquebeta@example.com" not in blob
        assert "gammaperson.deltaperson@example.com" not in blob
        assert "suppressed" in blob
        assert (
            test_client.patch(
                "/api/v1/organization/duty-activity-access-policy",
                json={"small_group_threshold": 2},
            ).status_code
            == 200
        )
        large = test_client.get(f"/api/v1/exports/duty-activity/works-council/{period_id}.xlsx")
        assert large.status_code == 200
        workbook = load_workbook(BytesIO(large.content))
        values = [
            str(value)
            for row in workbook.active.iter_rows(values_only=True)
            for value in row
            if value is not None
        ]
        blob = " ".join(values).lower()
        assert "uniquealpha" not in blob
        assert "gammaperson" not in blob
        assert "suppressed" not in blob
        assert "bereitschaftsdienst" in blob
        pdf = test_client.get(f"/api/v1/exports/duty-activity/works-council/{period_id}.pdf")
        assert pdf.status_code == 200
        pdf_text = pdf.content.lower()
        assert b"uniquealpha" not in pdf_text
        assert b"gammaperson" not in pdf_text


def test_purge_removes_only_expired_duty_activity_and_logs_run():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    testing_session = sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False)
    Base.metadata.create_all(engine)
    db = testing_session()
    db.add(Organization(id=1, name="Default", slug="default", plan_tier="team"))
    db.flush()
    member = TeamMember(
        organization_id=1,
        first_name="Keep",
        last_name="Name",
        email="keep@example.com",
        is_active=True,
    )
    db.add(member)
    db.flush()
    db.add(
        TimeEntry(
            organization_id=1,
            team_member_id=member.id,
            entry_date=date(2024, 9, 10),
            kind="call_out",
            source="duty_activity",
            duration_minutes=60,
        )
    )
    db.add(
        TimeEntry(
            organization_id=1,
            team_member_id=member.id,
            entry_date=date(2024, 9, 11),
            kind="in_duty_activity",
            source="duty_activity",
            duration_minutes=60,
        )
    )
    db.add(
        TimeEntry(
            organization_id=1,
            team_member_id=member.id,
            entry_date=date(2024, 9, 1),
            kind="work",
            source="roster",
            duration_minutes=480,
        )
    )
    db.commit()
    deleted = purge_expired_duty_activity_episodes(
        db,
        organization_id=1,
        actor="test",
        source="test",
        as_of=date(2026, 9, 11),
    )
    assert deleted == 1
    remaining = list(db.scalars(select(TimeEntry)))
    kinds_dates = {(row.kind, row.entry_date) for row in remaining}
    assert ("call_out", date(2024, 9, 10)) not in kinds_dates
    assert ("in_duty_activity", date(2024, 9, 11)) in kinds_dates
    assert ("work", date(2024, 9, 1)) in kinds_dates
    logs = list(
        db.scalars(
            select(AuditLog).where(AuditLog.entity_type == "duty_activity", AuditLog.action == "purge")
        )
    )
    assert len(logs) == 1
    assert logs[0].details["deleted"] == 1
    assert logs[0].details["cutoff"] == "2024-09-11"
    again = purge_expired_duty_activity_episodes(
        db,
        organization_id=1,
        actor="test",
        source="test",
        as_of=date(2026, 9, 11),
    )
    assert again == 0
    db.close()
    engine.dispose()
